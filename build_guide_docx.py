# -*- coding: utf-8 -*-
"""把《看板使用说明.md》导出为《看板使用说明.docx》。

排版：documents 技能预设 compact_reference_guide（US Letter、1in 页边距、9360 DXA 内容宽度、
正文 11pt/1.25 行距、表格固定列宽），仅把预设的蓝色标题改为看板品牌茶绿（brand_accent）。

用法：
    python -X utf8 build_guide_docx.py              # 只生成 DOCX
    python -X utf8 build_guide_docx.py --preview    # 额外输出 _preview.html（版式自查用）

改了 看板使用说明.md 之后重新跑一次即可，两个文件内容保持一致。"""
import io, re, html as H
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD = "看板使用说明.md"
OUT_DOCX = "看板使用说明.docx"
OUT_HTML = "_preview.html"

LATIN, EA, MONOF = "Calibri", "微软雅黑", "Consolas"
H1C, H2C, H3C = "45684A", "45684A", "2F4A33"
INK, SUB, HEADFILL, GRID = "1A1A1A", "6E6B66", "F1F0EC", "C9C6C0"
NOTE_FILL, MONO_FILL = "F7F6F3", "F5F4F1"
W = 9360
THREE = [2160, 3600, 3600]
LABEL_S, LABEL_L = [1701, 7659], [2700, 6660]

# ------------------------------------------------------------------ helpers
def E(tag, **kw):
    if ":" not in tag:
        tag = "w:" + tag
    e = OxmlElement(tag)
    for k, v in kw.items():
        e.set(qn("w:" + k), str(v))
    return e

RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
             "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
             "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs",
             "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
             "em", "lang", "eastAsianLayout", "specVanish", "oMath"]

def insert_ordered(parent, child, order):
    tag = child.tag.split("}")[1]
    idx = order.index(tag)
    for existing in parent:
        et = existing.tag.split("}")[1]
        if et in order and order.index(et) > idx:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child

def set_rfonts(rPr, latin, ea):
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = E("rFonts"); rPr.insert(0, rf)
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        if rf.get(qn("w:" + a)) is not None:
            del rf.attrib[qn("w:" + a)]
    rf.set(qn("w:ascii"), latin); rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), ea); rf.set(qn("w:cs"), latin)
    rf.set(qn("w:hint"), "eastAsia")

def run(p, text, size=11, bold=False, color=INK, mono=False):
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    set_rfonts(r._element.get_or_add_rPr(), MONOF if mono else LATIN, EA)
    return r

def shade(elm_getter, fill):
    pr = elm_getter()
    pr.append(E("shd", val="clear", color="auto", fill=fill))

def p_border_bottom(p, color=GRID, sz=6):
    pPr = p._p.get_or_add_pPr()
    bdr = E("pBdr"); bdr.append(E("bottom", val="single", sz=sz, space="1", color=color))
    pPr.append(bdr)

# --------------------------------------------------------------- md parsing
TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")

def smart(t):
    return re.sub(r'"([^"\n]*)"', "\u201c\\1\u201d", t)

def inline(text, bold=False):
    out = []
    for part in TOKEN.split(text):
        if part == "":
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.extend(inline(part[2:-2], True))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            opt = {"m": True}
            if bold:
                opt["b"] = True
            out.append((part[1:-1], opt))
        else:
            opt = {"b": True} if bold else {}
            out.append((smart(part), opt))
    return out

def dedent(ls):
    out = []
    for ln in ls:
        out.append(ln[3:] if ln.startswith("   ") else (ln[1:] if ln.startswith("\t") else ln))
    return out

def region_end(ls, start):
    j = start
    while j < len(ls):
        ln = ls[j]
        if ln.strip() == "":
            j += 1; continue
        if re.match(r"^\d+\.\s", ln):
            break
        if ln.startswith("   ") or ln.startswith("\t"):
            j += 1; continue
        break
    return j

def parse(ls):
    blocks, i = [], 0
    while i < len(ls):
        ln = ls[i]
        if ln.startswith("```"):
            i += 1; buf = []
            while i < len(ls) and not ls[i].startswith("```"):
                buf.append(ls[i]); i += 1
            i += 1; blocks.append(("mono", buf)); continue
        if ln.strip() == "---":
            i += 1; blocks.append(("hr",)); continue
        if ln.startswith(">"):
            buf = []
            while i < len(ls) and ls[i].startswith(">"):
                buf.append(ls[i].lstrip(">").strip()); i += 1
            blocks.append(("note", buf)); continue
        m = re.match(r"^(#+)\s+(.*)$", ln)
        if m:
            blocks.append(("h" + str(len(m.group(1))), m.group(2).strip())); i += 1; continue
        if ln.lstrip().startswith("|"):
            buf = []
            while i < len(ls) and ls[i].lstrip().startswith("|"):
                buf.append(ls[i].strip()); i += 1
            rows = []
            for r in buf:
                cells = [c.strip() for c in r.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    continue
                rows.append(cells)
            blocks.append(("tbl", rows)); continue
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", ln)
        if m and m.group(1) == "":
            items = []
            while i < len(ls):
                m2 = re.match(r"^(\d+)\.\s+(.*)$", ls[i])
                if not m2:
                    break
                text = m2.group(2); i += 1
                j = region_end(ls, i)
                sub = parse(dedent(ls[i:j])) if j > i else []
                items.append((text, sub)); i = j
            blocks.append(("ol", items)); continue
        m = re.match(r"^(\s*)[-*]\s+(.*)$", ln)
        if m:
            buf = []
            while i < len(ls):
                m2 = re.match(r"^(\s*)[-*]\s+(.*)$", ls[i])
                if not m2:
                    break
                buf.append((1 if len(m2.group(1)) >= 2 else 0, m2.group(2))); i += 1
            blocks.append(("ul", buf)); continue
        if ln.strip() == "":
            i += 1; continue
        blocks.append(("p", ln.strip())); i += 1
    return blocks

# ------------------------------------------------------------- docx building
def setup_numbering(doc):
    numbering = doc.part.numbering_part.element
    abs_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    cur = [max(abs_ids or [-1]), max(num_ids or [0])]

    def lvl(ilvl, fmt, text, font=None, left=540, hang=270):
        e = E("lvl", ilvl=ilvl)
        e.append(E("start", val=1)); e.append(E("numFmt", val=fmt)); e.append(E("lvlText", val=text))
        e.append(E("lvlJc", val="left"))
        pp = E("pPr"); pp.append(E("ind", left=left, hanging=hang)); e.append(pp)
        if font:
            rp = E("rPr"); rp.append(E("rFonts", ascii=font, hAnsi=font, hint="default")); e.append(rp)
        return e

    def add_abs(levels):
        cur[0] += 1
        a = E("abstractNum", abstractNumId=cur[0])
        a.append(E("multiLevelType", val="hybridMultilevel"))
        for l in levels:
            a.append(l)
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(a)
        else:
            numbering.append(a)
        cur[1] += 1
        n = E("num", numId=cur[1]); n.append(E("abstractNumId", val=cur[0])); numbering.append(n)
        return cur[1]

    bullet = add_abs([lvl(0, "bullet", "\uf0b7", "Symbol"), lvl(1, "bullet", "\uf0b7", "Symbol", left=1080, hang=270)])
    numbers = add_abs([lvl(0, "decimal", "%1.")])
    return bullet, numbers, cur[0]

def list_style(doc, name, numid, ilvl, size=11, after=4):
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles["Normal"]
    ppr = st.element.get_or_add_pPr()
    npr = E("numPr"); npr.append(E("ilvl", val=ilvl)); npr.append(E("numId", val=numid))
    ppr.append(npr)
    pf = st.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(after); pf.line_spacing = 1.25
    st.font.size = Pt(size)
    set_rfonts(st.element.get_or_add_rPr(), LATIN, EA)
    return st

def style_table(t, widths, tbl_ind=120):
    tblPr = t._tbl.tblPr
    for c in list(tblPr):
        tblPr.remove(c)
    tblPr.append(E("tblW", w=sum(widths), type="dxa"))
    tblPr.append(E("tblInd", w=tbl_ind, type="dxa"))
    b = E("tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(E(edge, val="single", sz=4, space=0, color=GRID))
    tblPr.append(b)
    tblPr.append(E("tblLayout", type="fixed"))
    cm = E("tblCellMar")
    for side, w in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
        cm.append(E(side, w=w, type="dxa"))
    tblPr.append(cm)
    tblPr.append(E("tblLook", val="04A0", firstRow=1, lastRow=0, firstColumn=1, lastColumn=0, noHBand=0, noVBand=1))
    grid = t._tbl.find(qn("w:tblGrid"))
    for gc in list(grid):
        grid.remove(gc)
    for w in widths:
        grid.append(E("gridCol", w=w))
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            for old in tcPr.findall(qn("w:tcW")):
                tcPr.remove(old)
            tcPr.insert(0, E("tcW", w=w, type="dxa"))

def cell_fill(cell, text, header=False, size=10):
    if header:
        shade(lambda: cell._tc.get_or_add_tcPr(), HEADFILL)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    parts = str(text).split("<br>")
    for k, part in enumerate(parts):
        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0 if k == len(parts) - 1 else 3)
        pf.line_spacing = 1.15
        for txt, opt in inline(part):
            run(p, txt, size=size, bold=header or opt.get("b", False),
                color=INK, mono=opt.get("m", False))

def pick_widths(rows):
    cols = max(len(r) for r in rows)
    if cols >= 3:
        return THREE
    labels = [r[0] for r in rows]
    need = 0
    for lb in labels:
        plain = re.sub(r"[*`]", "", lb)
        pt = sum((1.0 if ord(ch) > 0x2E80 else 0.5) for ch in plain) * 10
        need = max(need, pt)
    need_dxa = 240 + 100 + need * 20
    return LABEL_S if need_dxa <= LABEL_S[0] else LABEL_L

def scale_widths(widths, total):
    s0 = sum(widths)
    out = [int(round(w * total / float(s0))) for w in widths]
    out[-1] += total - sum(out)
    return out

def add_table(doc, rows, widths, indent=0):
    n = len(widths)
    tbl_ind = 120 + indent
    if indent:
        widths = scale_widths(widths, W - tbl_ind)
    t = doc.add_table(rows=1, cols=n)
    t.autofit = False
    first = True
    for k, row in enumerate(rows):
        cells = t.rows[0].cells if first else t.add_row().cells
        first = False
        for ci in range(n):
            cell_fill(cells[ci], row[ci] if ci < len(row) else "", header=(k == 0))
    t.rows[0]._tr.get_or_add_trPr().append(E("tblHeader"))
    style_table(t, widths, tbl_ind)
    sep = doc.add_paragraph()
    pf = sep.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
    pPr = sep._p.get_or_add_pPr()
    rPr = E("rPr"); rPr.append(E("sz", val=12)); rPr.append(E("szCs", val=12))
    pPr.append(rPr)
    return t

def add_page_field(p):
    r = p.add_run()
    r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string(SUB)
    set_rfonts(r._element.get_or_add_rPr(), LATIN, EA)
    b = E("fldChar", fldCharType="begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
    en = E("fldChar", fldCharType="end")
    r._element.append(b); r._element.append(it); r._element.append(en)

def build_docx(blocks):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    set_rfonts(normal.element.get_or_add_rPr(), LATIN, EA)
    npf = normal.paragraph_format
    npf.space_before = Pt(0); npf.space_after = Pt(6); npf.line_spacing = 1.25

    for name, size, color, before, after in (("Heading 1", 16, H1C, 18, 10),
                                             ("Heading 2", 13, H2C, 14, 7),
                                             ("Heading 3", 12, H3C, 10, 5)):
        st = doc.styles[name]
        st.font.size = Pt(size); st.font.bold = True; st.font.italic = False
        st.font.color.rgb = RGBColor.from_string(color)
        set_rfonts(st.element.get_or_add_rPr(), LATIN, EA)
        pf = st.paragraph_format
        pf.space_before = Pt(before); pf.space_after = Pt(after)
        pf.line_spacing = 1.25; pf.keep_with_next = True

    bn, nn, abs_num = setup_numbering(doc)
    list_style(doc, "GuideBullet", bn, 0)
    list_style(doc, "GuideBullet2", bn, 1)
    list_style(doc, "GuideNumber", nn, 0)

    footer = sec.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    fp.paragraph_format.space_before = Pt(0); fp.paragraph_format.space_after = Pt(0)
    run(fp, "喜茶华南区运营数据看板 · 使用说明", size=8.5, color=SUB)
    run(fp, "\t", size=8.5, color=SUB)
    run(fp, "第 ", size=8.5, color=SUB)
    add_page_field(fp)
    run(fp, " 页", size=8.5, color=SUB)

    def fresh_num():
        numbering = doc.part.numbering_part.element
        ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        nid = max(ids) + 1
        n = E("num", numId=nid); n.append(E("abstractNumId", val=abs_num))
        numbering.append(n)
        return nid

    def set_numpr(p, nid, ilvl=0):
        pPr = p._p.get_or_add_pPr()
        npr = E("numPr"); npr.append(E("ilvl", val=ilvl)); npr.append(E("numId", val=nid))
        pPr.append(npr)
        return p

    def para(segments, size=11, before=0, after=6, line=1.25, style=None, ind=None, color=INK, align=None):
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        pf = p.paragraph_format
        if style is None:
            pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
        if ind is not None:
            pf.left_indent = Inches(ind)
        if align is not None:
            p.alignment = align
        for txt, opt in segments:
            run(p, txt, size=opt.get("s", size), bold=opt.get("b", False),
                color=opt.get("c", color), mono=opt.get("m", False))
        return p

    # ---- title block (memo_masthead) ----
    title = blocks[0][1]
    meta = blocks[1][1] if len(blocks) > 1 and blocks[1][0] == "note" else []
    kick = doc.add_paragraph()
    kick.paragraph_format.space_before = Pt(0); kick.paragraph_format.space_after = Pt(2)
    kr = run(kick, "内部使用 · 华南战区", size=9, color=SUB)
    insert_ordered(kr._element.get_or_add_rPr(), E("spacing", val=30), RPR_ORDER)
    t = para([(title, {})], size=20, after=4)
    for r in t.runs:
        r.bold = True
    para([("三个视图怎么用 · 指标口径怎么算 · 数据怎么更新", {})], size=11.5, after=12, color=SUB)
    for line in meta:
        lab, _, val = line.partition("：")
        m = doc.add_paragraph()
        m.paragraph_format.space_before = Pt(0); m.paragraph_format.space_after = Pt(2)
        m.paragraph_format.line_spacing = 1.15
        run(m, lab + "：", size=10.5, bold=True, color=INK)
        run(m, val, size=10.5, color=INK)
    rule = doc.add_paragraph()
    p_border_bottom(rule)
    rule.paragraph_format.space_before = Pt(6); rule.paragraph_format.space_after = Pt(12)

    # ---- body ----
    def emit(blks, depth=0):
        for blk in blks:
            kind = blk[0]
            if kind == "hr":
                continue
            if kind == "h1":
                h = doc.add_heading(smart(blk[1]), level=1)
                h.paragraph_format.space_before = Pt(18); h.paragraph_format.space_after = Pt(10)
            elif kind == "h2":
                h = doc.add_heading(smart(blk[1]), level=1)
                h.paragraph_format.space_before = Pt(18); h.paragraph_format.space_after = Pt(10)
            elif kind == "h3":
                h = doc.add_heading(smart(blk[1]), level=2)
                h.paragraph_format.space_before = Pt(14); h.paragraph_format.space_after = Pt(7)
            elif kind == "p":
                txt = blk[1]
                is_q = re.sub(r"[*`]", "", txt).startswith("Q：")
                segs = inline(txt)
                plain = re.sub(r"[*`]", "", txt)
                if is_q:
                    for sg in segs:
                        sg[1]["b"] = True
                    para(segs, before=8, after=2)
                elif plain.startswith("内部运营工具"):
                    para(segs, size=9, before=12, after=0, color=SUB)
                else:
                    para(segs, after=8 if depth else 6)
            elif kind == "ul":
                for lvl, txt in blk[1]:
                    p = doc.add_paragraph(style="GuideBullet2" if (lvl or depth) else "GuideBullet")
                    for txt2, opt in inline(txt):
                        run(p, txt2, size=opt.get("s", 11), bold=opt.get("b", False),
                            color=opt.get("c", INK), mono=opt.get("m", False))
            elif kind == "ol":
                nid = fresh_num()
                for txt, subs in blk[1]:
                    p = set_numpr(doc.add_paragraph(style="GuideNumber"), nid)
                    for txt2, opt in inline(txt):
                        run(p, txt2, size=opt.get("s", 11), bold=opt.get("b", False),
                            color=opt.get("c", INK), mono=opt.get("m", False))
                    if subs:
                        emit(subs, depth + 1)
            elif kind == "note":
                body = " ".join(blk[1])
                p = doc.add_paragraph()
                shade(lambda: p._p.get_or_add_pPr(), NOTE_FILL)
                pf = p.paragraph_format
                pf.space_before = Pt(4); pf.space_after = Pt(8); pf.line_spacing = 1.25
                pf.left_indent = Inches(0.15)
                for txt, opt in inline(body):
                    run(p, txt, size=opt.get("s", 10.5), bold=opt.get("b", False),
                        color=opt.get("c", INK), mono=opt.get("m", False))
            elif kind == "mono":
                lines = blk[1]
                for k, ln in enumerate(lines):
                    p = doc.add_paragraph()
                    shade(lambda: p._p.get_or_add_pPr(), MONO_FILL)
                    pf = p.paragraph_format
                    pf.space_before = Pt(4 if k == 0 else 0)
                    pf.space_after = Pt(4 if k == len(lines) - 1 else 0)
                    pf.line_spacing = 1.15
                    pf.left_indent = Inches(0.15)
                    run(p, ln, size=9.5, mono=True)
            elif kind == "tbl":
                add_table(doc, blk[1], pick_widths(blk[1]), indent=540 if depth else 0)
    emit(blocks[2:] if meta else blocks[1:])

    cp = doc.core_properties
    cp.title = title
    cp.author = "喜茶 华南战区"
    cp.last_modified_by = "喜茶 华南战区"
    doc.save(OUT_DOCX)
    return title, meta

# ------------------------------------------------------------ html preview
def esc(s):
    return H.escape(s)

def il_html(text):
    out = []
    for txt, opt in inline(text):
        t = esc(txt)
        if opt.get("b"):
            t = "<b>" + t + "</b>"
        if opt.get("m"):
            t = "<code>" + t + "</code>"
        out.append(t)
    return "".join(out)

def build_html(blocks, title, meta):
    css = """
@page { size: 8.5in 11in; margin: 1in; }
* { box-sizing: border-box; }
body { margin: 0; background: #DDD; font-family: "微软雅黑", Calibri, sans-serif; color: #1A1A1A; }
.page { width: 8.5in; min-height: 11in; padding: 1in; margin: 0 auto 12px; background: #fff; }
.kicker { font-size: 9pt; letter-spacing: 1.5pt; color: #6E6B66; margin: 0 0 2pt; }
h1.title { font-size: 20pt; margin: 0 0 4pt; }
.sub { font-size: 11.5pt; color: #6E6B66; margin: 0 0 12pt; }
.meta { font-size: 10.5pt; line-height: 1.15; margin: 0 0 2pt; }
.rule { border-bottom: 0.75pt solid #C9C6C0; margin: 6pt 0 12pt; }
h2, h3 { color: #45684A; font-size: 16pt; margin: 18pt 0 10pt; page-break-after: avoid; }
h3 { font-size: 13pt; margin: 14pt 0 7pt; }
p { font-size: 11pt; line-height: 1.25; margin: 0 0 6pt; }
ul, ol { margin: 0 0 6pt; padding-left: 0.375in; }
li { font-size: 11pt; line-height: 1.25; margin-bottom: 4pt; }
table { border-collapse: collapse; table-layout: fixed; margin: 4pt 0 10pt 0.083in; }
th, td { border: 0.5pt solid #C9C6C0; padding: 4pt 6pt; font-size: 10pt; line-height: 1.15; vertical-align: top; overflow-wrap: anywhere; }
th { background: #F1F0EC; font-weight: bold; text-align: left; }
.note { background: #F7F6F3; font-size: 10.5pt; padding: 4pt 6pt; margin: 4pt 0 8pt 0.15in; }
pre { background: #F5F4F1; font-family: Consolas, monospace; font-size: 9.5pt; line-height: 1.15;
      padding: 4pt 6pt; margin: 4pt 0 8pt 0.15in; white-space: pre-wrap; }
code { font-family: Consolas, monospace; font-size: 9.5pt; }
.foot { font-size: 8.5pt; color: #6E6B66; border-top: 0.5pt solid #E8E6E1; padding-top: 4pt; }
"""
    o = [f'<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>{esc(title)}</title><style>{css}</style></head><body><div class="page">']
    o.append(f'<p class="kicker">内部使用 · 华南战区</p><h1 class="title">{esc(title)}</h1>')
    o.append('<p class="sub">三个视图怎么用 · 指标口径怎么算 · 数据怎么更新</p>')
    for line in meta:
        lab, _, val = line.partition("：")
        o.append(f'<p class="meta"><b>{esc(lab)}：</b>{esc(val)}</p>')
    o.append('<div class="rule"></div>')

    def emit(blks, depth=0):
        for blk in blks:
            k = blk[0]
            if k == "hr":
                continue
            if k in ("h1", "h2"):
                o.append(f"<h2>{esc(smart(blk[1]))}</h2>")
            elif k == "h3":
                o.append(f"<h3>{esc(smart(blk[1]))}</h3>")
            elif k == "p":
                txt = blk[1]
                is_q = re.sub(r"[*`]", "", txt).startswith("Q：")
                body = il_html(re.sub(r"^\*\*|\*\*$", "", txt) if is_q else txt)
                if is_q:
                    body = "<b>" + il_html(txt.replace("**", "")) + "</b>"
                cls = ' class="foot"' if re.sub(r"[*`]", "", txt).startswith("内部运营工具") else ""
                o.append(f'<p{cls} style="margin-top:{"8" if is_q else "0"}pt">{body}</p>')
            elif k == "ul":
                o.append("<ul>")
                nested = False
                for lvl, txt in blk[1]:
                    if lvl and not nested:
                        o.append('<ul style="margin-top:2pt">'); nested = True
                    elif not lvl and nested:
                        o.append("</ul>"); nested = False
                    o.append(f"<li>{il_html(txt)}</li>")
                if nested:
                    o.append("</ul>")
                o.append("</ul>")
            elif k == "ol":
                o.append("<ol>")
                for txt, subs in blk[1]:
                    o.append(f"<li>{il_html(txt)}")
                    if subs:
                        o.append('<div style="margin-top:2pt">')
                        emit(subs, depth + 1)
                        o.append("</div>")
                    o.append("</li>")
                o.append("</ol>")
            elif k == "note":
                body = " ".join(blk[1]).replace("**", "")
                o.append(f'<p class="note">{il_html(body)}</p>')
            elif k == "mono":
                o.append("<pre>" + esc("\n".join(blk[1])) + "</pre>")
            elif k == "tbl":
                rows = blk[1]
                widths = pick_widths(rows)
                o.append("<table><colgroup>" + "".join(f'<col style="width:{w/20/72:.3f}in">' for w in widths) + "</colgroup>")
                for ri, row in enumerate(rows):
                    tag = "th" if ri == 0 else "td"
                    o.append("<tr>")
                    for ci in range(len(widths)):
                        c = esc(row[ci]) if ci < len(row) else ""
                        inner = "<br>".join(il_html(x) for x in (row[ci] if ci < len(row) else "").split("<br>"))
                        o.append(f"<{tag}>{inner}</{tag}>")
                    o.append("</tr>")
                o.append("</table>")
    emit(blocks[2:] if meta else blocks[1:])
    o.append("</div></body></html>")
    io.open(OUT_HTML, "w", encoding="utf-8").write("\n".join(o))

def main():
    import sys
    blocks = parse(io.open(MD, encoding="utf-8").read().split("\n"))
    title, meta = build_docx(blocks)
    print("blocks:", len(blocks), "| meta rows:", len(meta))
    print("DOCX ->", OUT_DOCX)
    if "--preview" in sys.argv:
        build_html(blocks, title, meta)
        print("HTML preview ->", OUT_HTML)

if __name__ == "__main__":
    main()
